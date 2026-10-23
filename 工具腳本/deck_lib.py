# -*- coding: utf-8 -*-
"""自然色系簡報的版面元件，樣式從 2026-08-21 那 32 頁的 pptx XML 量出來。

這支檔案存在的理由：C2C-132 的建構腳本住在一個被清掉的暫存目錄裡，
下一輪要做簡報時整套工具都不見了，只剩 CLAUDE.md 對它的描述。
所以這一次把它放進 repo。

兩個設計決定，兩者都是 C2C-132 量到的缺陷換來的：

1. kv()／band_note()／card() 回傳它自己的結束位置。
   C2C-132 有 18 處色塊壓住表格，而 18 處全部是同一個原因：
   用手算的常數去放表格後面的元素。表格知道自己畫到哪裡，問它就好。

2. Deck.check() 是 build 期的版面守衛，不是事後檢查。
   它記下每個 kv 表格的真實 bbox，任何之後畫的東西與它相交、
   越出版面、或壓到頁碼，就讓 build 失敗。
   C2C-117 的重疊檢查只比「有文字的框」彼此，所以一個後畫的、
   沒有文字的裝飾色塊壓在文字上，它永遠回報乾淨。
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# 量出來的樣式（來源：文件/SROS2智慧防火牆_專題最終報告_自然色系_2026-08-21.pptx）
W_IN, H_IN = 13.3333333, 7.5
FONT = "Microsoft JhengHei"

INK   = RGBColor(0x2F, 0x30, 0x2D)   # 153 次，最常見的字色
PAPER = RGBColor(0xFA, 0xF8, 0xF3)   # 130 次
MUTED = RGBColor(0x66, 0x66, 0x5F)   #  98 次
SAGE  = RGBColor(0x65, 0x70, 0x65)   #  37 次填色
BRICK = RGBColor(0x86, 0x60, 0x5F)   #  32 次
TAN   = RGBColor(0xA1, 0x8F, 0x83)   #  18 次
STONE = RGBColor(0xD1, 0xCC, 0xC4)   #  64 次，最常見的填色
MIST  = RGBColor(0xEA, 0xE6, 0xDF)
BLUE  = RGBColor(0xBB, 0xC1, 0xCA)
GREY  = RGBColor(0x99, 0x98, 0x8D)

# 頁碼的位置也是量出來的：28／32 頁都在這裡
PAGENUM = (12.330, 7.030, 0.600, 0.220)

L0 = 0.670            # 左邊界
EYEBROW_T = 0.820
TITLE_T   = 1.140


def _char_w(ch):
    return 0.52 if ord(ch) < 0x2E80 else 1.0


def est_width(text, pt):
    """估計一行文字的寬度（吋）。這是字寬估計不是 render。"""
    return sum(_char_w(c) for c in text) * pt / 72.0


def est_lines(text, box_w, pt):
    """估計這段文字要幾行。"""
    total = 0
    for para in text.split(chr(10)):
        if not para.strip():
            total += 1
            continue
        w = est_width(para, pt)
        n = int(w / box_w) + (1 if w % box_w > 1e-9 else 0)
        total += max(1, n)
    return total


def _runs_text(runs):
    if isinstance(runs, str):
        return runs
    return "".join(r if isinstance(r, str) else r[0] for r in runs)


def fit_h(runs, box_w, pt, line=1.45, pad=0.06):
    """這段文字在這個寬度下需要多高。"""
    return est_lines(_runs_text(runs), box_w, pt) * pt * line / 72.0 + pad


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(W_IN)
        self.prs.slide_height = Inches(H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.tables = {}      # slide_idx -> [(L,T,W,H), ...]
        self.order = {}       # slide_idx -> 畫到第幾個
        self.violations = []

    # ---- 版面守衛 ----------------------------------------------------
    def _register(self, s, left, top, w, h, is_table=False, has_text=False,
                  exempt=False):
        i = list(self.prs.slides).index(s)
        n = self.order.get(i, 0)
        self.order[i] = n + 1
        eps = 1e-6

        # (a) 越界
        if left < -eps or top < -eps or left + w > W_IN + eps or top + h > H_IN + eps:
            self.violations.append(
                "p%d #%d 越界 L%.3f T%.3f W%.3f H%.3f" % (i + 1, n, left, top, w, h))

        # (b) 壓到之前登記的表格（C2C-132 那 18 處的形狀）
        if not is_table:
            for (tl, tt, tw, th) in self.tables.get(i, []):
                if (left < tl + tw - eps and left + w > tl + eps and
                        top < tt + th - eps and top + h > tt + eps):
                    self.violations.append(
                        "p%d #%d 壓住表格 元素T%.3f-%.3f 表格T%.3f-%.3f"
                        % (i + 1, n, top, top + h, tt, tt + th))

        # (c) 有文字的框壓到頁碼（頁碼自己當然不算）
        if has_text and not exempt:
            pl, pt_, pw, ph = PAGENUM
            if (left < pl + pw - eps and left + w > pl + eps and
                    top < pt_ + ph - eps and top + h > pt_ + eps):
                self.violations.append("p%d #%d 壓住頁碼" % (i + 1, n))

        if is_table:
            self.tables.setdefault(i, []).append((left, top, w, h))

    def check(self):
        if self.violations:
            raise SystemExit("版面守衛擋下 %d 處：\n  " % len(self.violations)
                             + "\n  ".join(self.violations))

    def save(self, path):
        self.check()
        self.prs.save(path)

    # ---- 元件 --------------------------------------------------------
    def slide(self, band=None):
        """band: None / (rail, color) / (panel, w, color) / (full, l, w, color)"""
        s = self.prs.slides.add_slide(self.blank)
        self.rect(s, 0, 0, W_IN, H_IN, PAPER)
        if band is None:
            pass
        elif band[0] == "rail":
            self.rect(s, 0, 0, 0.150, H_IN, band[1])
        elif band[0] == "panel":
            self.rect(s, W_IN - band[1], 0, band[1], H_IN, band[2])
        elif band[0] == "full":
            self.rect(s, band[1], 0, band[2], H_IN, band[3])
        return s

    def rect(self, s, left, top, w, h, fill):
        sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                Inches(w), Inches(h))
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
        sh.line.fill.background()
        sh.shadow.inherit = False
        sh.text_frame.word_wrap = True
        self._register(s, left, top, w, h)
        return sh

    def text(self, s, left, top, w, h, runs, size=13.5, color=INK,
             align=PP_ALIGN.LEFT, space=0.0, bold=False, line=1.35,
             exempt=False):
        """runs: str 或 [(text, color|None, bold|None), ...]；換行用 \\n 分段。"""
        tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if isinstance(runs, str):
            runs = [(runs, None, None)]
        # 允許清單裡夾雜裸字串：("abc") 是 str 不是 tuple，而那個括號很容易看成 tuple
        runs = [(r, None, None) if isinstance(r, str) else r for r in runs]
        paras = [[]]
        for t, c, b in runs:
            parts = t.split("\n")
            for k, part in enumerate(parts):
                if k:
                    paras.append([])
                if part:
                    paras[-1].append((part, c, b))
        for pi, pr in enumerate(paras):
            p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            for t, c, b in pr:
                r = p.add_run()
                r.text = t
                r.font.name = FONT
                r.font.size = Pt(size)
                r.font.bold = bold if b is None else b
                r.font.color.rgb = color if c is None else c
                if space:
                    r.font._rPr.set("spc", str(int(space * 100)))
        self._register(s, left, top, w, h, has_text=True, exempt=exempt)
        return tb

    def eyebrow(self, s, t, left=L0, w=6.0):
        self.text(s, left, EYEBROW_T, w, 0.260, t, size=10.5, color=TAN, space=1.2)

    def title(self, s, t, w=11.2, size=28.5, min_h=0.760):
        h = max(min_h, fit_h(t, w, size, line=1.26))
        self.text(s, L0, TITLE_T, w, h, t, size=size, color=INK, bold=True, line=1.26)
        return TITLE_T + h

    def lead(self, s, t, top, w=11.2, min_h=0.500, size=13.5):
        h = max(min_h, fit_h(t, w, size))
        self.text(s, L0, top, w, h, t, size=size, color=MUTED, line=1.45)
        return top + h

    def rule(self, s, top, w, left=L0, color=TAN, h=0.022):
        self.rect(s, left, top, w, h, color)
        return top + h

    def kv(self, s, rows, top, left=L0, w=8.3, header=None, gap=0.330,
           kw=0.45, size=13.5, vcolor=None):
        """回傳表格的結束 y —— 永遠用它，不要手算。"""
        y = top
        kwid = w * kw - 0.02
        vl = left + w * kw
        vwid = w * (1 - kw)
        if header:
            self.text(s, left, y, kwid, 0.240, header[0], size=8.6, color=TAN, space=1.0)
            self.text(s, vl, y, vwid, 0.240, header[1], size=8.6, color=TAN, space=1.0)
            y += 0.300
            self.rect(s, left, y, w, 0.014, MIST)
            y += 0.140
        for k, v in rows:
            kr = k if isinstance(k, list) else [(k, None, None)]
            vr = v if isinstance(v, list) else [(v, vcolor, None)]
            rh = max(gap, fit_h(kr, kwid, size, line=1.30, pad=0.02),
                     fit_h(vr, vwid, size, line=1.30, pad=0.02))
            self.text(s, left, y, kwid, rh, kr, size=size, color=INK, line=1.30)
            self.text(s, vl, y, vwid, rh, vr, size=size, color=INK, line=1.30)
            y += rh
        self._register(s, left, top, w, y - top, is_table=True)
        return y

    def band_note(self, s, top, runs, left=L0, w=8.3, h=0.860, fill=MIST, size=10.5):
        h = max(h, fit_h(runs, w - 0.44, size) + 0.30)
        self.rect(s, left, top, w, h, fill)
        self.text(s, left + 0.22, top + 0.15, w - 0.44, h - 0.30, runs,
                  size=size, color=INK, line=1.45)
        return top + h

    def card(self, s, left, top, label, value, fill=STONE, w=2.900, h=1.180,
             vcolor=None, vsize=20.2):
        self.rect(s, left, top, w, h, fill)
        tc = PAPER if fill in (BRICK, SAGE) else TAN
        vc = (PAPER if fill in (BRICK, SAGE) else INK) if vcolor is None else vcolor
        self.text(s, left + 0.17, top + 0.14, w - 0.34, 0.240, label,
                  size=8.6, color=tc, space=1.0)
        self.text(s, left + 0.17, top + 0.42, w - 0.34, 0.620, value,
                  size=vsize, color=vc, bold=True)
        return top + h

    def footer(self, s, runs, top=6.980, left=L0, w=9.0, size=8.2):
        self.text(s, left, top, w, 0.330, runs, size=size, color=GREY, line=1.40)

    def pagenum(self, s, n):
        l, t, w, h = PAGENUM
        self.text(s, l, t, w, h, "%02d" % n, size=8.2, color=TAN,
                  align=PP_ALIGN.RIGHT, exempt=True)
